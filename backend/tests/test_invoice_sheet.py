from datetime import datetime

import pytest
from fastapi import HTTPException
from openpyxl import Workbook

from backend.app.main import order_invoice, product_catalog
from backend.app.services.invoice_sheet_service import get_order_invoice_data, list_catalog_products


def _workbook(tmp_path):
    path = tmp_path / "demo_process_kpi_data.xlsx"
    workbook = Workbook()
    invoices = workbook.active
    invoices.title = "Invoice_date"
    invoices.append(["Order ID", "Invoice ID", "Company", "Customer", "Order Date", "Billing Address", "Shipping Address", "HSN", "Contact Details", None])
    invoices.append(["ORD-1", "INV-1", "Northstar Health", "Emily Carter", datetime(2026, 10, 1), "1 Bill St, Dallas, TX 75230", "2 Ship St, Dallas, TX 75230", 902131, "emily@example.com", None])
    invoices.append(["ORD-2", None, "MedCare Systems", "Michael Turner", datetime(2026, 10, 2), "3 Main St, Atlanta, GA 30309", "3 Main St, Atlanta, GA 30309", 902131, "michael@example.com", None])
    orders = workbook.create_sheet("Order")
    orders.append(["Order ID", "Description", "Customer Name", "PO Number", "Invoice Status", "Invoice Amount"])
    orders.append(["ORD-1", "Existing PO invoice flow", "Northstar Health", "PO-1001", "Approved", 5000])
    orders.append(["ORD-2", "New order", "MedCare Systems", "None", "Pending", "Null"])
    orders.append(["ORD-3", "Order only", "City General", "PO-3", "Not Started", 100])
    products = workbook.create_sheet("Prod_cat")
    products.append(["Product_Id", "Product Category", "Product Family", "Product / Variant"])
    products.append(["PROD-001", "Breast Implant", "MemoryGel™ Smooth Round", "MemoryGel™ Breast Implants"])
    products.append([None, None, None, None])
    workbook.save(path)
    return path


def test_order_invoice_data_combines_invoice_and_order_sheets(tmp_path):
    path = _workbook(tmp_path)
    assert get_order_invoice_data("ord-1", path) == {
        "order_id": "ORD-1",
        "invoice_id": "INV-1",
        "company": "Northstar Health",
        "customer": "Emily Carter",
        "order_date": "2026-10-01",
        "billing_address": "1 Bill St, Dallas, TX 75230",
        "shipping_address": "2 Ship St, Dallas, TX 75230",
        "hts_code": "902131",
        "contact_details": "emily@example.com",
        "po_number": "PO-1001",
        "order_description": "Existing PO invoice flow",
        "invoice_status": "Approved",
        "invoice_amount": "5000",
    }
    pending = get_order_invoice_data("ORD-2", path)
    assert pending["invoice_id"] == "" and pending["po_number"] == "" and pending["invoice_amount"] == ""
    order_only = get_order_invoice_data("ORD-3", path)
    assert order_only["company"] == "City General" and order_only["billing_address"] == ""
    assert get_order_invoice_data("ORD-404", path) is None


def test_catalog_products_are_read_from_prod_cat(tmp_path):
    assert list_catalog_products(_workbook(tmp_path)) == [{
        "product_id": "PROD-001",
        "category": "Breast Implant",
        "family": "MemoryGel™ Smooth Round",
        "name": "MemoryGel™ Breast Implants",
    }]


def test_invoice_endpoints_serve_workbook_data():
    data = order_invoice("ORD-78421")
    assert data["company"] and data["billing_address"]
    assert product_catalog()
    with pytest.raises(HTTPException):
        order_invoice("ORD-DOES-NOT-EXIST")
