from __future__ import annotations

from datetime import date, datetime, time
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from backend.app.services.demo_dashboard_service import DEMO_WORKBOOK_PATH


FEDEX_SHEET = "FedEx_Shipments"


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == time() else value.isoformat(timespec="minutes")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    text = str(value).strip()
    return "" if text.casefold() in {"none", "null"} else text


def _keys(headers: tuple[Any, ...]) -> list[str]:
    """Snake-case headers; repeated headers (e.g. Commercial_Invoice_No) get a numeric suffix so no column is lost."""
    seen: dict[str, int] = {}
    keys = []
    for header in headers:
        key = _text(header).lower().replace(" ", "_")
        seen[key] = seen.get(key, 0) + 1
        keys.append(key if seen[key] == 1 else f"{key}_{seen[key]}")
    return keys


def list_fedex_shipments(
    order_id: str | None = None,
    workbook_path: Path = DEMO_WORKBOOK_PATH,
) -> list[dict[str, str]]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        if FEDEX_SHEET not in workbook.sheetnames:
            return []
        rows = workbook[FEDEX_SHEET].iter_rows(values_only=True)
        keys = _keys(next(rows, ()))
        shipments = [
            {key: _text(value) for key, value in zip(keys, row) if key}
            for row in rows
            if any(value is not None for value in row)
        ]
    finally:
        workbook.close()
    wanted = (order_id or "").strip().casefold()
    return [s for s in shipments if s.get("order_id")] if not wanted else [
        s for s in shipments if s.get("order_id", "").casefold() == wanted
    ]
