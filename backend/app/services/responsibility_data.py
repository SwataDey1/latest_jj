from __future__ import annotations

from collections import defaultdict
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


REQUIRED_SHEETS = ("Responsibilities", "Responsibility_Work", "Activities_Tasks")
RESPONSIBILITY_COLUMNS = {"responsibility_id", "persona_id", "responsibility", "metric_name"}
WORK_COLUMNS = {"responsibility_id", "work_item_id", "work_item", "status"}
ACTIVITY_COLUMNS = {"persona_id", "activity_name"}


class ResponsibilityDataError(ValueError):
    """Raised when the workbook cannot provide responsibility data."""


def _read_sheet(
    workbook_path: Path,
    sheet_name: str,
    required_columns: set[str],
) -> list[dict[str, Any]]:
    try:
        workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    except FileNotFoundError as error:
        raise ResponsibilityDataError(f"Workbook not found: {workbook_path}") from error
    try:
        if sheet_name not in workbook.sheetnames:
            raise ResponsibilityDataError(f"Missing required sheet: {sheet_name}")
        sheet = workbook[sheet_name]
        rows = sheet.iter_rows(values_only=True)
        try:
            headers = [str(value).strip() if value is not None else "" for value in next(rows)]
        except StopIteration as error:
            raise ResponsibilityDataError(f"Sheet {sheet_name} is empty") from error
        missing_columns = required_columns - set(headers)
        if missing_columns:
            missing = ", ".join(sorted(missing_columns))
            raise ResponsibilityDataError(f"Sheet {sheet_name} is missing columns: {missing}")
        records = [
            {headers[index]: value for index, value in enumerate(row) if index < len(headers)}
            for row in rows
            if any(value is not None and str(value).strip() for value in row)
        ]
        if not records:
            raise ResponsibilityDataError(f"Sheet {sheet_name} has no records")
        return records
    finally:
        workbook.close()


def _text(value: Any) -> str:
    return str(value).strip() if value is not None else ""


def _first_value(row: dict[str, Any], *names: str) -> str:
    for name in names:
        value = _text(row.get(name))
        if value:
            return value
    return ""


def _persona_from_responsibility_id(responsibility_id: str) -> str:
    parts = responsibility_id.split("-")
    return "-".join(parts[:2]) if len(parts) >= 2 else ""


def load_responsibility_data(
    persona_id: str,
    workbook_path: Path,
    responsibility_id: str | None = None,
    status: str | None = None,
) -> dict[str, Any]:
    responsibilities = _read_sheet(workbook_path, "Responsibilities", RESPONSIBILITY_COLUMNS)
    work_rows = _read_sheet(workbook_path, "Responsibility_Work", WORK_COLUMNS)
    activity_rows = _read_sheet(workbook_path, "Activities_Tasks", ACTIVITY_COLUMNS)

    scoped_responsibilities = [
        row for row in responsibilities
        if _text(row.get("persona_id")) == persona_id
        and _text(row.get("responsibility_id"))
        and (responsibility_id is None or _text(row.get("responsibility_id")) == responsibility_id)
    ]
    if responsibility_id and not scoped_responsibilities:
        return {"responsibilities": [], "status_filter": status, "data_source": workbook_path.name}

    activity_names = [
        _text(row.get("activity_name"))
        for row in activity_rows
        if _text(row.get("persona_id")) == persona_id and _text(row.get("activity_name"))
    ]
    work_by_responsibility: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in work_rows:
        current_responsibility_id = _text(row.get("responsibility_id"))
        row_persona = _text(row.get("persona_id")) or _persona_from_responsibility_id(current_responsibility_id)
        if row_persona == persona_id:
            work_by_responsibility[current_responsibility_id].append(row)

    selected_status = status.casefold() if status else None
    result: list[dict[str, Any]] = []
    for responsibility_index, responsibility in enumerate(scoped_responsibilities):
        current_responsibility_id = _text(responsibility.get("responsibility_id"))
        source_work = work_by_responsibility.get(current_responsibility_id, [])
        status_counts: dict[str, int] = defaultdict(int)
        for row in source_work:
            current_status = _text(row.get("status"))
            if current_status:
                status_counts[current_status] += 1
        linked_work = [
            row for row in source_work
            if selected_status is None or _text(row.get("status")).casefold() == selected_status
        ]
        normalized_work = []
        for work_index, row in enumerate(linked_work):
            activity_name = activity_names[(responsibility_index + work_index) % len(activity_names)] if activity_names else _text(responsibility.get("metric_name"))
            due_date = _first_value(row, "due_date", "due_at", "due")
            normalized_work.append({
                "work_item_id": _first_value(row, "work_item_id") or f"{current_responsibility_id}-work-{work_index + 1}",
                "responsibility_id": current_responsibility_id,
                "customer_request_id": _first_value(row, "customer_request_id"),
                "work_item": _first_value(row, "work_item", "work_item_name"),
                "work_item_name": _first_value(row, "work_item", "work_item_name"),
                "activity_name": activity_name,
                "status": _text(row.get("status")),
                "priority": _first_value(row, "priority", "urgency"),
                "owner": _first_value(row, "owner", "assigned_resource", "assigned_resource_name", "resource_name") or persona_id,
                "updated_at": _first_value(row, "updated_at", "updated_at_timestamp", "due_at") or due_date,
                "due_at": due_date,
                "due_date": due_date,
                "responsibility": _text(responsibility.get("responsibility")),
            })
        result.append({
            "responsibility_id": current_responsibility_id,
            "responsibility": _text(responsibility.get("responsibility")),
            "responsibility_name": _text(responsibility.get("responsibility")),
            "activity_name": activity_names[responsibility_index % len(activity_names)] if activity_names else _text(responsibility.get("metric_name")),
            "activities": activity_names,
            "metric_name": _text(responsibility.get("metric_name")),
            "formula": _text(responsibility.get("formula")),
            "target_area": _text(responsibility.get("target_area")),
            "status_counts": dict(status_counts),
            "work_items": normalized_work,
            "work_item_count": len(normalized_work),
        })
    return {
        "responsibilities": result,
        "status_filter": status,
        "data_source": workbook_path.name,
    }
