from __future__ import annotations

import csv
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from backend.app.config import ROOT_DIR


_WORKBOOK_PRIMARY = ROOT_DIR / "data" / "demo_process_kpi_data.xlsx"
_WORKBOOK_CANDIDATES = (
    _WORKBOOK_PRIMARY,
    *sorted((ROOT_DIR / "data").glob("demo_process_kpi_data *.xlsx"), reverse=True),
)
CUSTOMER_MASTER_WORKBOOK_PATH = next(
    (path for path in _WORKBOOK_CANDIDATES if path.exists()),
    _WORKBOOK_CANDIDATES[0],
)
CUSTOMER_MASTER_SHEET = "Customer_master"
CUSTOMER_APPLICATIONS_PATH = ROOT_DIR / "data" / "customer_applications.csv"


def _read_csv(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as source:
        return list(csv.DictReader(source))


def list_customers() -> list[dict[str, str]]:
    workbook = load_workbook(CUSTOMER_MASTER_WORKBOOK_PATH, read_only=True, data_only=True)
    try:
        sheet_name = next(
            name for name in workbook.sheetnames
            if name.casefold() == CUSTOMER_MASTER_SHEET.casefold()
        )
        rows = workbook[sheet_name].iter_rows(values_only=True)
        headers = next(rows)
        return [
            {
                str(header): "" if value is None else str(value)
                for header, value in zip(headers, row)
            }
            for row in rows
            if any(value is not None for value in row)
        ]
    finally:
        workbook.close()


def list_customer_applications() -> list[dict[str, str]]:
    return _read_csv(CUSTOMER_APPLICATIONS_PATH)


def find_customer(username: str) -> dict[str, str] | None:
    normalized_username = username.strip().casefold()
    return next(
        (
            customer
            for customer in list_customers()
            if customer.get("username", "").strip().casefold() == normalized_username
        ),
        None,
    )


def find_customer_by_id(customer_id: str | int) -> dict[str, str] | None:
    normalized_id = str(customer_id).strip()
    return next(
        (
            customer
            for customer in list_customers()
            if customer.get("customer_id", "").strip() == normalized_id
        ),
        None,
    )


def get_applications(customer_id: str | int) -> list[dict[str, Any]]:
    normalized_id = str(customer_id).strip()
    return [
        {
            "application_name": application["application_name"],
            "tenancy_id": application["tenancy_id"],
            "rights": [
                right.strip()
                for right in application.get("rights", "").split("|")
                if right.strip()
            ],
        }
        for application in list_customer_applications()
        if application.get("customer_id", "").strip() == normalized_id
    ]


def get_login_options() -> dict[str, list[str]]:
    customers = list_customers()
    return {
        "categories": sorted({customer["customer_category"] for customer in customers if customer.get("customer_category")}),
        "personas": sorted({customer["persona"] for customer in customers if customer.get("persona")}),
    }
