from __future__ import annotations

from datetime import date, datetime, time
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from backend.app.services.demo_dashboard_service import DEMO_WORKBOOK_PATH


# The invoice sheet has been named differently across workbook revisions.
INVOICE_SHEETS = ("Invoice_data", "Invoice_Data", "Invoice_date", "Invoice_Master")
ORDER_SHEET = "Order"
PRODUCT_SHEET = "Prod_cat"


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == time() else value.isoformat(timespec="minutes")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    text = str(value).strip()
    return "" if text.casefold() in {"none", "null"} else text


def _rows(workbook: Any, sheet_name: str) -> list[dict[str, str]]:
    if sheet_name not in workbook.sheetnames:
        return []
    rows = workbook[sheet_name].iter_rows(values_only=True)
    headers = [_text(header) for header in next(rows, ())]
    return [
        {header: _text(value) for header, value in zip(headers, row) if header}
        for row in rows
        if any(value is not None for value in row)
    ]


def _find(rows: list[dict[str, str]], order_id: str) -> dict[str, str] | None:
    wanted = order_id.strip().casefold()
    return next((row for row in rows if row.get("Order ID", "").casefold() == wanted), None)


def list_catalog_products(workbook_path: Path = DEMO_WORKBOOK_PATH) -> list[dict[str, str]]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        rows = _rows(workbook, PRODUCT_SHEET)
    finally:
        workbook.close()
    return [
        {
            "product_id": row.get("Product_Id", ""),
            "category": row.get("Product Category", ""),
            "family": row.get("Product Family", ""),
            "name": row.get("Product / Variant", ""),
        }
        for row in rows
        if row.get("Product_Id")
    ]


def get_order_invoice_data(order_id: str, workbook_path: Path = DEMO_WORKBOOK_PATH) -> dict[str, Any] | None:
    """Invoice details for an order from the invoice sheet, enriched with PO/amount data from the Order sheet."""
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        sheet = next((name for name in INVOICE_SHEETS if name in workbook.sheetnames), None)
        invoice = _find(_rows(workbook, sheet), order_id) if sheet else None
        order = _find(_rows(workbook, ORDER_SHEET), order_id)
    finally:
        workbook.close()
    if invoice is None and order is None:
        return None
    invoice = invoice or {}
    order = order or {}
    return {
        "order_id": invoice.get("Order ID") or order.get("Order ID", ""),
        "invoice_id": invoice.get("Invoice ID", ""),
        "company": invoice.get("Company") or order.get("Customer Name", ""),
        "customer": invoice.get("Customer", ""),
        "order_date": invoice.get("Order Date", ""),
        "billing_address": invoice.get("Billing Address", ""),
        "shipping_address": invoice.get("Shipping Address", ""),
        "hts_code": invoice.get("HSN", ""),
        "contact_details": invoice.get("Contact Details", ""),
        "po_number": order.get("PO Number", ""),
        "order_description": order.get("Description", ""),
        "invoice_status": order.get("Invoice Status", ""),
        "invoice_amount": order.get("Invoice Amount", ""),
    }
